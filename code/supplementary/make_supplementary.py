#!/usr/bin/env python3
"""Builds the Supplementary Tables S1-S3 of the article from the released records.

    python3 code/supplementary/make_supplementary.py [--out results/supplementary]

Writes one CSV per table and a workbook with one sheet per table
(Supplementary_Tables_S1-S3.xlsx). Python 3.10+, openpyxl for the workbook.
Before writing, it checks the counts that the article reports for Corpus B and Corpus A and stops if any differs.
"""
import argparse, collections, csv, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
from corpus_b_state import final_state  # noqa: E402
from columns import COLUMNS  # noqa: E402

V = lambda e: e.get("value") if isinstance(e, dict) and "value" in e else e


def j(v):
    """Lists as '; '-separated text; None as empty."""
    if v is None:
        return ""
    if isinstance(v, list):
        return "; ".join(str(x) for x in v)
    return str(v)


def bib_labels():
    txt = open(f"{ROOT}/results/corpus-b/corpus_b_2026-09-18.bib", encoding="utf-8").read()
    out = {}
    for ent in re.split(r"\n(?=@)", txt):
        k = re.search(r"note = \{Corpus B, (rayyan-\d+)\}", ent)
        if not k:
            continue
        g = lambda f: (re.search(rf"\b{f} = \{{(.*?)\}},?\n", ent, re.S) or [None, ""])[1]
        a = g("author").split(" and ")
        first = a[0].split(",")[0].strip()
        lab = f"{first} et al. ({g('year')})" if len(a) > 2 else (f"{first} and {a[1].split(',')[0].strip()} ({g('year')})" if len(a) == 2 else f"{first} ({g('year')})")
        out[k.group(1)] = dict(report=lab, year=g("year"), title=g("title").strip("{}"), doi=g("doi"))
    return out


# ---------------------------------------------------------------- S1
def table_s1():
    A = 'TITLE-ABS-KEY("ISO 50001" OR "ISO 50006") AND PUBYEAR > 1995 AND DOCTYPE(ar OR cp OR re)'
    B = ('TITLE-ABS-KEY(("energy performance indicator*" OR "EnPI" OR "energy baseline*" OR "EnB" OR '
         '"energy performance measurement" OR "energy performance monitoring") AND ("industrial" OR "industry" OR '
         '"manufacturing" OR "facility" OR "organization" OR "management system" OR "energy management" OR '
         '"ISO 50001" OR "continuous process" OR "refin*")) AND PUBYEAR > 1995 AND DOCTYPE(ar OR cp OR re)')
    def wos(q):
        q = q.replace("TITLE-ABS-KEY(", "TS=(").replace(" AND PUBYEAR > 1995", " AND PY=(1996-2026)")
        return q.replace(" AND DOCTYPE(ar OR cp OR re)", " AND DT=(Article OR Proceedings Paper OR Review)")
    note_s = "As registered in the protocol (OSF, DOI 10.17605/OSF.IO/TDW6C)."
    note_w = ("Scopus string with the substitutions stated in Appendix A of the article: TITLE-ABS-KEY by TS, "
              "PUBYEAR > 1995 by PY=(1996-2026), document types by DT=(Article OR Proceedings Paper OR Review).")
    hdr = ["Database", "String", "Query", "Records", "Searched", "Exported", "Note"]
    rows = [["Scopus", "A (normative)", A, 468, "February 2026", "24 February 2026", note_s],
            ["Scopus", "B (methodological)", B, 486, "February 2026", "24 February 2026", note_s],
            ["Web of Science Core Collection", "A (normative)", wos(A), 263, "February 2026", "24 February 2026", note_w],
            ["Web of Science Core Collection", "B (methodological)", wos(B), 261, "February 2026", "24 February 2026", note_w]]
    return hdr, rows


# ---------------------------------------------------------------- S2
def tables_s2(R, log, feed):
    lab = bib_labels()
    rep = lambda k: lab.get(k, {}).get("report", k)
    models, meas, mech, adp, art = [], [], [], [], []
    for k in sorted(R, key=lambda k: (rep(k), k)):
        r = R[k]
        nm = 0
        for x in r["applications"]:
            for m in x["models"]:
                nm += 1
                c1, c2, c4 = m["C1_dependent_variable"], m["C2_model_form"], m["C4_data"]
                c3 = m.get("C3_variables") or []
                types = {v["type"] for v in c3 if v["role"] in ("predictor", "denominator")}
                c5 = m.get("C5_measurement") or []
                oos = m["C8_metrics"]["out_of_sample_validation"]
                models.append([rep(k), k, x["application_id"], m["model_id"], V(x["sector_registered"]),
                               V(x["refinery_or_petrochemical_case"]), V(x["scale_of_application"]), V(x["country_region"]),
                               V(c1["quantity"]), V(c1["energy_vector"]), V(c1.get("denominator")),
                               j(V(c2["model_family"])), V(c2.get("estimation_method")), V(c2.get("n_independent_variables")),
                               "; ".join(f"{v['verbatim']} [{v['type']}, {v['role']}]" for v in c3),
                               "yes" if (types and types <= {"throughput"}) else "no",
                               V(c4.get("modelling_frequency")), V(c4.get("span_months")), V(c4.get("n_observations")), V(c4.get("data_type")),
                               len(c5), sum(1 for c in c5 if c["boundary_correspondence"] == "exact"),
                               V(m["C6_re_estimation"]["regime"]), V(m["C6_re_estimation"].get("update_status")),
                               V(m["C7_validity_envelope"]["statement"]), V(m["C7_validity_envelope"].get("bounds_verbatim")),
                               V(m["C8_metrics"]["in_sample_fit"]["status"]), V(oos["status"]), V(oos.get("validation_design"))])
                for c in c5:
                    meas.append([rep(k), k, x["application_id"], m["model_id"], c.get("energy_vector"), c.get("data_origin"),
                                 c.get("boundary_correspondence"), c.get("page"), c.get("support_class")])
        for x in r["mechanisms"]:
            mech.append([rep(k), k, x["mechanism_id"], x.get("refers_to"), x["claim_status"], x["mechanism_cause"],
                         x["mechanism_effect_on_model"], x.get("effect_other_verbatim"), x.get("page"), x.get("support_class")])
        for x in r["adaptations"]:
            adp.append([rep(k), k, x["adaptation_id"], x.get("refers_to"), x["adaptation"], x.get("adaptation_other_verbatim"),
                        x["status"], x.get("addresses_mechanism"), x["normative_framing"], x.get("page"), x.get("support_class")])
        c13 = r["C13_reproducibility"]
        art.append([rep(k), k, lab.get(k, {}).get("title", ""), lab.get(k, {}).get("doi", ""), nm, len(r["mechanisms"]), len(r["adaptations"]),
                    V(c13["equations_given"]), V(c13["coefficients_given"]), V(c13["data_available"]), V(c13["code_available"]),
                    "yes" if k in feed else "no"])
    for row in models + meas + mech + adp + art:
        for i, v in enumerate(row):
            row[i] = j(v) if isinstance(v, list) or v is None else v
    H = {
        "S2a_reports": ["Report", "Record key", "Title", "DOI", "Models", "Mechanism occurrences", "Adaptation occurrences",
                        "Equations given", "Coefficients given", "Data available", "Code available",
                        "Feed or product characteristics discussed, not used as variables"],
        "S2b_models": ["Report", "Record key", "Application", "Model", "Sector (registered list)", "Refinery or petrochemical context",
                       "Scale", "Country or region", "Dependent quantity", "Energy carrier", "Denominator", "Model family",
                       "Estimation method", "Independent variables (n)", "Variables [type, role]", "Load is the only normalising variable",
                       "Modelling frequency", "Span (months)", "Observations (n)", "Data type", "Measurement entries (n)",
                       "Entries with exact boundary correspondence (n)", "Re-estimation regime", "Update status",
                       "Validity domain", "Validity bounds (verbatim)", "In-sample fit", "Out-of-sample result", "Validation design"],
        "S2c_measurement": ["Report", "Record key", "Application", "Model", "Energy carrier", "Data origin", "Boundary correspondence", "Page", "Support"],
        "S2d_mechanisms": ["Report", "Record key", "Mechanism (record ID)", "Refers to", "Claim status", "Cause", "Effect on the model", "Effect (verbatim, if other)", "Page", "Support"],
        "S2e_adaptations": ["Report", "Record key", "Adaptation (record ID)", "Refers to", "Type", "Type (verbatim, if other)", "State", "Addresses mechanism", "Normative framing", "Page", "Support"],
        "S2f_corrections": ["Check", "Item", "Record key", "Report", "Location", "Field", "Before", "After"],
    }
    corr = [[l["check"], l["item"], l["key"], rep(l["key"]), l["path"], l["field"], j(l["before"]), j(l["after"])] for l in log]
    return H, {"S2a_reports": art, "S2b_models": models, "S2c_measurement": meas, "S2d_mechanisms": mech,
               "S2e_adaptations": adp, "S2f_corrections": corr}


# ---------------------------------------------------------------- S3
FIELDS = [("paper_type", "Study type"), ("multi_site", "Multi-site"), ("sector_of_activity", "Sector"),
          ("country_region", "Country or region"), ("type_of_organisation", "Organisation type"),
          ("ems_standard", "Energy-management standard"), ("mv_protocol", "M&V protocol"),
          ("regulatory_driver", "Regulatory driver"), ("enpi_enb_model_types", "Model families"),
          ("complementary_methodologies", "Complementary methods")]


def table_s3():
    A = list(csv.DictReader(open(f"{ROOT}/results/corpus-a/tables/analysis_dataset.csv", encoding="utf-8")))
    st = {(r["paper_id"], r["field"]): r["status"] for r in
          csv.DictReader(open(f"{ROOT}/data/04-extraction-corpus-a/reconciliation/cell_provenance.csv", encoding="utf-8"))}
    def val(v):
        v = v.strip()
        if v.startswith("["):
            try:
                return "; ".join(json.loads(v))
            except json.JSONDecodeError:
                return v
        return v
    def authors(v):
        try:
            a = json.loads(v)
            return "; ".join(a) if isinstance(a, list) else v
        except json.JSONDecodeError:
            return v
    hdr = ["Record key", "Authors", "Year", "Title", "Source", "DOI", "Document type", "In Corpus B"]
    for _, lab in FIELDS:
        hdr += [lab, f"{lab}: status"]
    rows = []
    for r in sorted(A, key=lambda r: (r["year"], r["paper_id"])):
        row = [r["paper_id"], authors(r["authors"]), r["year"], r["title"], r["source"], r["doi"], r["document_type"],
               "yes" if r["in_corpus_b"] == "True" else "no"]
        for f, _ in FIELDS:
            row += [val(r[f]), st[(r["paper_id"], f)]]
        rows.append(row)
    return hdr, rows


# ---------------------------------------------------------------- checks against the article
def check(R, s3):
    models = [(k, m) for k, r in R.items() for x in r["applications"] for m in x["models"]]
    c5 = [(k, c) for k, m in models for c in m.get("C5_measurement") or []]
    cnt = lambda f: collections.Counter(f(m) for _, m in models)
    regime = cnt(lambda m: V(m["C6_re_estimation"]["regime"]))
    dom = cnt(lambda m: V(m["C7_validity_envelope"]["statement"]))
    oos = [k for k, m in models if V(m["C8_metrics"]["out_of_sample_validation"]["status"]) == "reported"]
    mech = [x for r in R.values() for x in r["mechanisms"]]
    claim = collections.Counter(x["claim_status"] for x in mech)
    c13 = lambda f, v: sum(1 for r in R.values() if V(r["C13_reproducibility"][f]) == v)
    only = sum(1 for _, m in models if (t := {v["type"] for v in m.get("C3_variables") or [] if v["role"] in ("predictor", "denominator")}) and t <= {"throughput"})
    got = {
        "reports": len(R), "models": len(models), "measurement entries": len(c5),
        "exact boundary entries": sum(1 for _, c in c5 if c["boundary_correspondence"] == "exact"),
        "articles with exact boundary": len({k for k, c in c5 if c["boundary_correspondence"] == "exact"}),
        "load the only normalising variable": only,
        "declared re-estimation regime": sum(n for v, n in regime.items() if v not in ("fixed, not mentioned", "not stated")),
        "event-triggered regime": regime["event-triggered, trigger stated"], "regime not stated": regime["not stated"],
        "validity domain, numerical": dom["stated with numeric bounds"], "validity domain, qualitative": dom["stated qualitatively"],
        "out-of-sample models": len(oos), "out-of-sample articles": len(set(oos)),
        "in-sample fit reported": cnt(lambda m: V(m["C8_metrics"]["in_sample_fit"]["status"]))["reported"],
        "mechanism occurrences": len(mech), "general": claim["general"], "applied": claim["applied to the case"],
        "demonstrated": claim["empirically demonstrated"], "adaptations": sum(len(r["adaptations"]) for r in R.values()),
        "equations yes": c13("equations_given", "yes"), "equations partly": c13("equations_given", "partial"),
        "data yes": c13("data_available", "yes"), "data partly": c13("data_available", "partial"),
        "code yes": c13("code_available", "yes"), "code partly": c13("code_available", "partial"),
        "Corpus A publications": len(s3), "Corpus A in Corpus B": sum(1 for r in s3 if r[7] == "yes"),
        "adaptations linked to a mechanism": sum(1 for r in R.values() for x in r["adaptations"] if x.get("addresses_mechanism")),
        "links to a missing mechanism": sum(1 for r in R.values() for x in r["adaptations"]
                                            if x.get("addresses_mechanism") and x["addresses_mechanism"] not in {m["mechanism_id"] for m in r["mechanisms"]}),
    }
    want = {"reports": 39, "models": 140, "measurement entries": 163, "exact boundary entries": 35, "articles with exact boundary": 11,
            "load the only normalising variable": 78, "declared re-estimation regime": 24, "event-triggered regime": 17, "regime not stated": 27,
            "validity domain, numerical": 13, "validity domain, qualitative": 6, "out-of-sample models": 15, "out-of-sample articles": 5,
            "in-sample fit reported": 49, "mechanism occurrences": 130, "general": 51, "applied": 52, "demonstrated": 27, "adaptations": 88,
            "equations yes": 24, "equations partly": 9, "data yes": 4, "data partly": 4, "code yes": 0, "code partly": 1,
            "Corpus A publications": 331, "Corpus A in Corpus B": 40,
            "adaptations linked to a mechanism": 41, "links to a missing mechanism": 0}
    bad = {k: (got[k], want[k]) for k in want if got[k] != want[k]}
    if bad:
        sys.exit(f"counts differ from the article: {bad}")
    print(f"checked {len(want)} counts against the article: all equal")


def write(out, sheets):
    os.makedirs(out, exist_ok=True)
    for name, (hdr, rows) in sheets.items():
        with open(f"{out}/{name}.csv", "w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh); w.writerow(hdr); w.writerows(rows)
    with open(f"{out}/columns.csv", "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh); w.writerow(["Table", "Column", "Meaning", "Values"]); w.writerows(COLUMNS)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed: CSV files only"); return
    wb = Workbook(); ws = wb.active; ws.title = "Notes"
    notes = [
        "Supplementary Tables S1-S3",
        "Ribeiro, B.; Miranda, N.; Eusébio, M. Energy Baseline Validity in Continuous Process Industries: A Scoping Review and Implications for Emission-Reduction Estimates. Submitted to Advances in Carbon Neutrality, 2026.",
        "",
        "S1: search strings and record counts by database.",
        "S2a-S2f: Corpus B charting records (39 charted reports) in the state analysed in the article: the extracted records, with the decisions of the September 2026 model verification and the corrections of the later checks listed in S2f. Verbatim quotations and pages for every value are in the deposited records (data/05-extraction-corpus-b/records/).",
        "S3: the 331 Corpus A publications with the field values on which the two model extractions agreed; the status column of each field says whether the value is a model agreement (used in the article), a single extraction, an unresolved divergence or missing (not used).",
        "Columns sheet: the meaning of every column and of the codes and values (A1 and M1, application and model; K1, mechanism; D1, adaptation within a report, unrelated to the deviations D1-D17 of the article; and so on).",
        "",
        "Values were proposed by language models and adjudicated or verified as described in the article (Sections 3.4, 3.5 and 4.8); they are not independently double-checked by humans except where the article says so.",
        "Generated by code/supplementary/make_supplementary.py in the review materials (Data Availability Statement).",
    ]
    for i, t in enumerate(notes, 1):
        ws.cell(i, 1, t).font = Font(name="Arial", size=11 if i == 1 else 10, bold=(i == 1))
    ws.column_dimensions["A"].width = 140
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wc = wb.create_sheet("Columns")
    wc.append(["Table", "Column", "Meaning", "Values"])
    for r in COLUMNS:
        wc.append(list(r))
    for c in wc[1]:
        c.font = Font(name="Arial", size=10, bold=True); c.fill = PatternFill("solid", fgColor="E8E6F5")
    for row in wc.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCD", (7, 34, 70, 90)):
        wc.column_dimensions[col].width = w
    wc.freeze_panes = "A2"
    titles = {"S1_search_strings": "S1", "S2a_reports": "S2a", "S2b_models": "S2b", "S2c_measurement": "S2c",
              "S2d_mechanisms": "S2d", "S2e_adaptations": "S2e", "S2f_corrections": "S2f", "S3_corpus_a": "S3"}
    hfill = PatternFill("solid", fgColor="E8E6F5")
    for name, (hdr, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(hdr)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(name="Arial", size=10, bold=True); c.fill = hfill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial", size=10)
        for i, h in enumerate(hdr, 1):
            L = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:400]])
            ws.column_dimensions[get_column_letter(i)].width = max(8, min(60, L + 2))
            if L > 60:
                for c in ws[get_column_letter(i)][1:]:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2" if name != "S1_search_strings" else "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(f"{out}/Supplementary_Tables_S1-S3.xlsx")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", default=f"{ROOT}/results/supplementary"); a = ap.parse_args()
    R, _, log, feed = final_state(ROOT)
    h3, s3 = table_s3()
    check(R, s3)
    sheets = {"S1_search_strings": table_s1()}
    H, T = tables_s2(R, log, feed)
    for n in H:
        sheets[n] = (H[n], T[n])
    sheets["S3_corpus_a"] = (h3, s3)
    write(a.out, sheets)
    for n, (h, r) in sheets.items():
        print(f"{n}: {len(r)} rows")


if __name__ == "__main__":
    main()
